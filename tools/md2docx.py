# -*- coding: utf-8 -*-
"""Markdown -> docx（面向表格密集的本体驱动需求规格说明书）
- 标题 -> Word 标题样式（导航窗格可用）；表格 -> 真 Word 表格（首行跨页重复）
- 代码块（```text / ```mermaid）-> 宋体等宽块，**按块自动适配字号**保证不折行
- 附录 D（22 屏 ASCII UI 原型）自动切横向分节，正文保持纵向
用法：python md2docx.py <源.md> <目标.docx>
"""
import io, re, sys, unicodedata
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, DST = sys.argv[1], sys.argv[2]
MARGIN_X, MARGIN_Y = 1.5, 1.6
PORT_W, PORT_H = 21.0, 29.7
LAND_W, LAND_H = 29.7, 21.0
LANDSCAPE_FROM = "附录 D UI 原型"
PORTRAIT_FROM = "需求完整性自检报告"
FONT = "宋体"

def set_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    rf.set(qn('w:ascii'), name); rf.set(qn('w:hAnsi'), name); rf.set(qn('w:eastAsia'), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = color

def shade(paragraph, fill="F2F2F2"):
    pr = paragraph._p.get_or_add_pPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    pr.append(sh)

def tight(paragraph, before=0, after=0, line=None):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    if line: pf.line_spacing = line

def repeat_header(row):
    tr = row._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); tr.append(th)

def width_units(s):
    return sum(2 if unicodedata.east_asian_width(c) in ("W", "F") else 1 for c in s)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")

def add_inline(paragraph, text, base_size=10.5):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            set_font(paragraph.add_run(tok[2:-2]), FONT, base_size, bold=True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1])
            set_font(r, FONT, base_size - 1.5)
            r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        else:
            set_font(paragraph.add_run(tok), FONT, base_size)

md = io.open(SRC, encoding="utf-8").read().splitlines()
doc = Document()
landscape = [False]

def setup_section(s, is_land):
    if is_land:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(LAND_W), Cm(LAND_H)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(PORT_W), Cm(PORT_H)
    s.left_margin = s.right_margin = Cm(MARGIN_X)
    s.top_margin = s.bottom_margin = Cm(MARGIN_Y)
    return s

setup_section(doc.sections[0], False)
st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def switch(land):
    if landscape[0] == land:
        return
    landscape[0] = land
    setup_section(doc.add_section(WD_SECTION.NEW_PAGE), land)

def text_width_pt(is_land):
    return (LAND_W if is_land else PORT_W) * 28.3465 - 2 * MARGIN_X * 28.3465

H = {1: (16, '1F3864'), 2: (14, '2E5496'), 3: (12, '2E5496'), 4: (11, '333333')}
i, n = 0, len(md)
tables = codes = heads = 0

while i < n:
    line = md[i]
    if line.startswith("```"):
        lang = line[3:].strip()
        buf, i = [], i + 1
        while i < n and not md[i].startswith("```"):
            buf.append(md[i]); i += 1
        i += 1
        maxw = max((width_units(b) for b in buf), default=1) or 1
        # 宋体排版宽度：半角(单位1)≈0.55×字号，全角(单位2)≈1.1×字号
        avail = text_width_pt(landscape[0]) * 0.96
        size = max(6.0, min(8.0, (avail / (maxw * 0.55)) // 0.5 * 0.5))
        if lang:
            cap = doc.add_paragraph(); tight(cap, 2, 0)
            set_font(cap.add_run(f"[{lang}]"), FONT, 7.5, italic=True, color=RGBColor(0x80, 0x80, 0x80))
        for b in buf:
            p = doc.add_paragraph(); tight(p, 0, 0, 1.0); shade(p)
            set_font(p.add_run(b if b.strip() else " "), FONT, size)
        codes += 1
        doc.add_paragraph()
        continue
    if line.startswith("|"):
        rows = []
        while i < n and md[i].startswith("|"):
            if not re.match(r"^\|[\s:\-|]+\|$", md[i].strip()):
                rows.append([c.strip() for c in md[i].strip().strip("|").split("|")])
            i += 1
        if rows:
            ncol = max(len(r) for r in rows)
            rows = [r + [""] * (ncol - len(r)) for r in rows]
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = True
            tw = (LAND_W if landscape[0] else PORT_W) - 2 * MARGIN_X
            first = tw * 1.35 / (ncol + 0.35)
            other = (tw - first) / max(ncol - 1, 1)
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    c = t.cell(ri, ci); c.width = Cm(first if ci == 0 else other)
                    p = c.paragraphs[0]; tight(p, 0, 0)
                    add_inline(p, cell, base_size=8.5)
                    if ri == 0:
                        for r in p.runs: r.font.bold = True
            repeat_header(t.rows[0])
            tables += 1
            doc.add_paragraph()
        continue
    m = re.match(r"^(#{1,6})\s+(.*)$", line)
    if m:
        title = m.group(2)
        if title.startswith(LANDSCAPE_FROM):
            switch(True)
        elif title.startswith(PORTRAIT_FROM):
            switch(False)
        lvl = min(len(m.group(1)), 4)
        size, color = H[lvl]
        p = doc.add_paragraph(style=f"Heading {lvl}")
        tight(p, 10 if lvl <= 2 else 6, 4)
        set_font(p.add_run(title), "微软雅黑", size, bold=True, color=RGBColor.from_string(color))
        heads += 1
        i += 1
        continue
    if line.startswith(">"):
        p = doc.add_paragraph(); tight(p, 2, 2)
        p.paragraph_format.left_indent = Cm(0.5)
        add_inline(p, line.lstrip("> ").rstrip(), base_size=9.5)
        for r in p.runs:
            r.font.italic = True; r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        i += 1
        continue
    m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
    if m:
        p = doc.add_paragraph(style='List Bullet' if len(m.group(1)) < 2 else 'List Bullet 2')
        tight(p, 0, 0); add_inline(p, m.group(2), base_size=10)
        i += 1
        continue
    m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
    if m:
        p = doc.add_paragraph(style='List Number'); tight(p, 0, 0)
        add_inline(p, m.group(2), base_size=10)
        i += 1
        continue
    if re.match(r"^\s*(---|===)\s*$", line):
        p = doc.add_paragraph(); tight(p, 2, 2)
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
        bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '6'); bot.set(qn('w:color'), 'BFBFBF')
        pbdr.append(bot); pPr.append(pbdr)
        i += 1
        continue
    if not line.strip():
        i += 1
        continue
    p = doc.add_paragraph(); tight(p, 0, 2)
    add_inline(p, line, base_size=10.5)
    i += 1

doc.save(DST)
print(f"标题 {heads} | 表格 {tables} | 代码块 {codes} | 分节 {len(doc.sections)} -> {DST}")
